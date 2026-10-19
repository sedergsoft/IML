from pptx import Presentation
from pptx.util import Inches, Pt
from pathlib import Path

from pathlib import Path

TEMPLATE = Path(r"D:\maestria\I semestre\Введение в машинное обучение\lab Регрессионные модели диагностика\IML\laboratorio_1\variante_13\lab_template.pptx")
OUT      = Path(r"D:\maestria\I semestre\Введение в машинное обучение\lab Регрессионные модели диагностика\IML\laboratorio_1\variante_13\laboratorio_1_variante_13.pptx")
FIG      = Path(r"D:\maestria\I semestre\Введение в машинное обучение\lab Регрессионные модели диагностика\IML\laboratorio_1\variante_13\figures")

prs = Presentation(TEMPLATE)

def set_text(slide, title, body_lines):
    # Asume que la diapositiva tiene placeholders de título y contenido
    title_ph = slide.shapes.title
    if title_ph is not None:
        title_ph.text = title
    body_ph = None
    for shp in slide.placeholders:
        if shp.placeholder_format.idx == 1:
            body_ph = shp
            break
    if body_ph is None:
        return
    tf = body_ph.text_frame
    tf.clear()
    for i, line in enumerate(body_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)

def add_picture(slide, filename, left, top, width):
    path = FIG / filename
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))

# --- Diapositiva 1: portada ---
s = prs.slides[0]
set_text(s, "Data learning", [
    "Curso \"Машинное обучение\" — Лабораторная работа",
    "Лаборатория анализа данных",
    "",
    "Tema: Regresión lineal y detección de valores atípicos (Variante 1-13)",
    "",
    "[Apellido] I.O., [Grupo], Variante 1-13",
])

# --- Diapositiva 2: datos ---
s = prs.slides[1]
set_text(s, "Datos de partida", [
    "Conjunto de datos: 110 observaciones, variables x (entrada) e y (salida).",
    "División holdout 70/30 (random_state=42): 77 entrenamiento, 33 prueba.",
    "La relación es aproximadamente lineal, con un grupo de puntos extremos en x≈3–4.5, y≈7–8.5.",
])
add_picture(s, "01_regresion_comparada.png", 5.5, 1.2, 4.0)

# --- Diapositiva 3: métodos ---
s = prs.slides[2]
set_text(s, "Métodos y fórmulas utilizadas", [
    "Modelo: y = β0 + β1·x + ε",
    "OLS: β1 = Σ(xi−x̄)(yi−ȳ) / Σ(xi−x̄)² ; β0 = ȳ − β1·x̄",
    "R² = 1 − SSE/SST",
    "Residuos estandarizados: ri = ei / (σ̂·√(1−hi))",
    "Leverage: hi (diagonal de la matriz sombrero)",
    "Distancia de Cook: Di = (ei² / (p·σ̂²)) · hi / (1−hi)²",
    "Umbral de filtrado: Di > 3·media(D)",
])

# --- Diapositiva 4: resultados ---
s = prs.slides[3]
set_text(s, "Resultados de la investigación", [
    "Modelo OLS inicial: ŷ = −0.311776 + 0.936361·x",
    "R² entrenamiento = 0.290399 ; R² prueba = −0.935444",
    "Máx |residuo estandarizado| = 2.513206 ; máx leverage = 0.097313 ; máx Cook = 0.125413",
    "Umbral de Cook = 0.053161 → 10 observaciones excluidas (n entrenamiento = 67)",
    "Modelo filtrado: ŷ = 0.465597 + 0.186327·x",
    "R² entrenamiento = 0.067465 ; R² prueba = −0.376683",
])
add_picture(s, "01_regresion_comparada.png", 0.3, 3.8, 3.2)
add_picture(s, "08_cook_vs_x.png", 3.6, 3.8, 3.2)
add_picture(s, "10_boxplot_comparacion.png", 6.9, 3.8, 3.2)

# --- Diapositiva 5: conclusiones ---
s = prs.slides[4]
set_text(s, "Conclusiones", [
    "1. La regresión OLS inicial explica solo una fracción limitada de la variabilidad y generaliza mal (R² prueba negativo).",
    "2. Residuos estandarizados, leverage y distancia de Cook son diagnósticos complementarios; ningún indicador por sí solo basta.",
    "3. Al filtrar las 10 observaciones con Cook por encima del umbral, la pendiente cambia de 0.936 a 0.186: esos puntos influían fuertemente en la recta.",
    "4. El R² de prueba mejora de −0.935 a −0.377, aunque sigue siendo insatisfactorio.",
    "5. No conviene eliminar observaciones automáticamente: verificar si son errores o casos legítimos; explorar modelos no lineales o robustos.",
])

prs.save(OUT)
print(f"Guardado: {OUT}")