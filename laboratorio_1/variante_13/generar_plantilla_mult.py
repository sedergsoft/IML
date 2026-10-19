from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

ROOT     = Path(__file__).resolve().parent
TEMPLATE = ROOT / "lab_template.pptx"
FIG      = ROOT / "figures"

# ---------------------------------------------------------------------------
# CONTENIDO POR IDIOMA
# ---------------------------------------------------------------------------
CONTENT = {
    "ru": {
        "suffix": "ru",
        "slides": [
            {
                "title": "Data learning",
                "body": [
                    "Курс «Машинное обучение» — Лабораторная работа",
                    "Лаборатория анализа данных",
                    "",
                    "Тема: Линейная регрессия и обнаружение выбросов (Вариант 1-13)",
                    "",
                    "[Фамилия] И.О., [Группа], Вариант 1-13",
                ],
                "pictures": [],
            },
            {
                "title": "Исходные данные",
                "body": [
                    "Набор данных: 110 наблюдений, переменные x (вход) и y (выход).",
                    "Разбиение holdout 70/30 (random_state=42): 77 обучающих, 33 тестовых.",
                    "Связь приблизительно линейная, но есть группа экстремальных точек при x≈3–4.5, y≈7–8.5.",
                ],
                "pictures": [("01_regresion_comparada.png", 5.5, 1.2, 4.0)],
            },
            {
                "title": "Используемые методы и формулы",
                "body": [
                    "Модель: y = β0 + β1·x + ε",
                    "OLS: β1 = Σ(xi−x̄)(yi−ȳ) / Σ(xi−x̄)² ; β0 = ȳ − β1·x̄",
                    "R² = 1 − SSE/SST",
                    "Стандартизованные остатки: ri = ei / (σ̂·√(1−hi))",
                    "Leverage: hi (диагональ hat-матрицы)",
                    "Расстояние Кука: Di = (ei² / (p·σ̂²)) · hi / (1−hi)²",
                    "Порог фильтрации: Di > 3·среднее(D)",
                ],
                "pictures": [],
            },
            {
                "title": "Результаты исследований",
                "body": [
                    "Начальная модель OLS: ŷ = −0.311776 + 0.936361·x",
                    "R² обучение = 0.290399 ; R² тест = −0.935444",
                    "Макс |станд. остаток| = 2.513206 ; макс leverage = 0.097313 ; макс Cook = 0.125413",
                    "Порог Кука = 0.053161 → исключено 10 наблюдений (n обучение = 67)",
                    "Отфильтрованная модель: ŷ = 0.465597 + 0.186327·x",
                    "R² обучение = 0.067465 ; R² тест = −0.376683",
                ],
                "pictures": [
                    ("01_regresion_comparada.png", 0.3, 3.8, 3.2),
                    ("08_cook_vs_x.png",            3.6, 3.8, 3.2),
                    ("10_boxplot_comparacion.png",  6.9, 3.8, 3.2),
                ],
            },
            {
                "title": "Выводы",
                "body": [
                    "1. Начальная OLS-регрессия объясняет лишь ограниченную часть вариации и плохо обобщает (R² тест отрицателен).",
                    "2. Стандартизованные остатки, leverage и расстояние Кука — взаимодополняющие диагностики; ни один индикатор сам по себе не достаточен.",
                    "3. После удаления 10 наблюдений с Cook выше порога наклон меняется с 0.936 на 0.186: эти точки сильно влияли на прямую.",
                    "4. R² тест улучшается с −0.935 до −0.377, но остаётся неудовлетворительным.",
                    "5. Не следует автоматически удалять наблюдения: нужно проверить, ошибки ли это или законные случаи; рекомендуется исследовать нелинейные или робастные модели.",
                ],
                "pictures": [],
            },
        ],
    },
    "es": {
        "suffix": "es",
        "slides": [
            {
                "title": "Data learning",
                "body": [
                    "Curso «Машинное обучение» — Лабораторная работа",
                    "Лаборатория анализа данных",
                    "",
                    "Tema: Regresión lineal y detección de valores atípicos (Variante 1-13)",
                    "",
                    "[Apellido] I.O., [Grupo], Variante 1-13",
                ],
                "pictures": [],
            },
            {
                "title": "Datos de partida",
                "body": [
                    "Conjunto de datos: 110 observaciones, variables x (entrada) e y (salida).",
                    "División holdout 70/30 (random_state=42): 77 entrenamiento, 33 prueba.",
                    "La relación es aproximadamente lineal, con un grupo de puntos extremos en x≈3–4.5, y≈7–8.5.",
                ],
                "pictures": [("01_regresion_comparada.png", 5.5, 1.2, 4.0)],
            },
            {
                "title": "Métodos y fórmulas utilizadas",
                "body": [
                    "Modelo: y = β0 + β1·x + ε",
                    "OLS: β1 = Σ(xi−x̄)(yi−ȳ) / Σ(xi−x̄)² ; β0 = ȳ − β1·x̄",
                    "R² = 1 − SSE/SST",
                    "Residuos estandarizados: ri = ei / (σ̂·√(1−hi))",
                    "Leverage: hi (diagonal de la matriz sombrero)",
                    "Distancia de Cook: Di = (ei² / (p·σ̂²)) · hi / (1−hi)²",
                    "Umbral de filtrado: Di > 3·media(D)",
                ],
                "pictures": [],
            },
            {
                "title": "Resultados de la investigación",
                "body": [
                    "Modelo OLS inicial: ŷ = −0.311776 + 0.936361·x",
                    "R² entrenamiento = 0.290399 ; R² prueba = −0.935444",
                    "Máx |residuo estandarizado| = 2.513206 ; máx leverage = 0.097313 ; máx Cook = 0.125413",
                    "Umbral de Cook = 0.053161 → 10 observaciones excluidas (n entrenamiento = 67)",
                    "Modelo filtrado: ŷ = 0.465597 + 0.186327·x",
                    "R² entrenamiento = 0.067465 ; R² prueba = −0.376683",
                ],
                "pictures": [
                    ("01_regresion_comparada.png", 0.3, 3.8, 3.2),
                    ("08_cook_vs_x.png",            3.6, 3.8, 3.2),
                    ("10_boxplot_comparacion.png",  6.9, 3.8, 3.2),
                ],
            },
            {
                "title": "Conclusiones",
                "body": [
                    "1. La regresión OLS inicial explica solo una fracción limitada de la variabilidad y generaliza mal (R² prueba negativo).",
                    "2. Residuos estandarizados, leverage y distancia de Cook son diagnósticos complementarios; ningún indicador por sí solo basta.",
                    "3. Al filtrar las 10 observaciones con Cook por encima del umbral, la pendiente cambia de 0.936 a 0.186: esos puntos influían fuertemente en la recta.",
                    "4. El R² de prueba mejora de −0.935 a −0.377, aunque sigue siendo insatisfactorio.",
                    "5. No conviene eliminar observaciones automáticamente: verificar si son errores o casos legítimos; explorar modelos no lineales o robustos.",
                ],
                "pictures": [],
            },
        ],
    },
    "en": {
        "suffix": "en",
        "slides": [
            {
                "title": "Data learning",
                "body": [
                    "Course “Machine Learning” — Laboratory work",
                    "Data Analysis Laboratory",
                    "",
                    "Topic: Linear regression and outlier detection (Variant 1-13)",
                    "",
                    "[Surname] I.O., [Group], Variant 1-13",
                ],
                "pictures": [],
            },
            {
                "title": "Initial data",
                "body": [
                    "Dataset: 110 observations, variables x (input) and y (output).",
                    "Holdout split 70/30 (random_state=42): 77 training, 33 test.",
                    "The relationship is roughly linear, with a group of extreme points at x≈3–4.5, y≈7–8.5.",
                ],
                "pictures": [("01_regresion_comparada.png", 5.5, 1.2, 4.0)],
            },
            {
                "title": "Methods and formulas used",
                "body": [
                    "Model: y = β0 + β1·x + ε",
                    "OLS: β1 = Σ(xi−x̄)(yi−ȳ) / Σ(xi−x̄)² ; β0 = ȳ − β1·x̄",
                    "R² = 1 − SSE/SST",
                    "Standardized residuals: ri = ei / (σ̂·√(1−hi))",
                    "Leverage: hi (diagonal of the hat matrix)",
                    "Cook's distance: Di = (ei² / (p·σ̂²)) · hi / (1−hi)²",
                    "Filtering threshold: Di > 3·mean(D)",
                ],
                "pictures": [],
            },
            {
                "title": "Research results",
                "body": [
                    "Initial OLS model: ŷ = −0.311776 + 0.936361·x",
                    "R² train = 0.290399 ; R² test = −0.935444",
                    "Max |standardized residual| = 2.513206 ; max leverage = 0.097313 ; max Cook = 0.125413",
                    "Cook threshold = 0.053161 → 10 observations excluded (n train = 67)",
                    "Filtered model: ŷ = 0.465597 + 0.186327·x",
                    "R² train = 0.067465 ; R² test = −0.376683",
                ],
                "pictures": [
                    ("01_regresion_comparada.png", 0.3, 3.8, 3.2),
                    ("08_cook_vs_x.png",            3.6, 3.8, 3.2),
                    ("10_boxplot_comparacion.png",  6.9, 3.8, 3.2),
                ],
            },
            {
                "title": "Conclusions",
                "body": [
                    "1. The initial OLS regression explains only a limited fraction of the variability and generalizes poorly (negative test R²).",
                    "2. Standardized residuals, leverage and Cook's distance are complementary diagnostics; no single indicator is sufficient.",
                    "3. After removing the 10 observations with Cook above the threshold, the slope changes from 0.936 to 0.186: those points strongly influenced the line.",
                    "4. The test R² improves from −0.935 to −0.377, but remains unsatisfactory.",
                    "5. Observations should not be removed automatically: check whether they are measurement errors or legitimate cases; explore nonlinear or robust models.",
                ],
                "pictures": [],
            },
        ],
    },
}

# ---------------------------------------------------------------------------
# ESCRITURA ROBUSTA DE TEXTO (usa placeholders si existen; si no, crea cuadros)
# ---------------------------------------------------------------------------
def fill_slide(slide, title, body_lines):
    used_body = False

    title_ph = slide.shapes.title
    if title_ph is not None:
        title_ph.text = title
    else:
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.9))
        tb.text_frame.text = title
        tb.text_frame.paragraphs[0].font.size = Pt(32)
        tb.text_frame.paragraphs[0].font.bold = True

    body_ph = None
    for shp in slide.placeholders:
        if shp.placeholder_format.idx != 0 and shp != title_ph and shp.has_text_frame:
            body_ph = shp
            break

    if body_ph is not None:
        tf = body_ph.text_frame
        tf.clear()
        for i, line in enumerate(body_lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(16)
        used_body = True

    if not used_body:
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(9), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(body_lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(16)

def add_pictures(slide, pictures):
    for fname, left, top, width in pictures:
        path = FIG / fname
        if path.exists():
            slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))
        else:
            print(f"  [aviso] no encontrada: {path}")

# ---------------------------------------------------------------------------
# GENERACIÓN
# ---------------------------------------------------------------------------
def build(lang, data):
    prs = Presentation(TEMPLATE)
    n_slides = len(prs.slides)
    for idx, spec in enumerate(data["slides"]):
        if idx >= n_slides:
            print(f"  [aviso] {lang}: la plantilla tiene {n_slides} diapositivas; se omite la {idx+1}")
            break
        slide = prs.slides[idx]
        fill_slide(slide, spec["title"], spec["body"])
        add_pictures(slide, spec["pictures"])
    out = ROOT / f"laboratorio_1_variante_13_{data['suffix']}.pptx"
    prs.save(out)
    print(f"Guardado: {out}")

if __name__ == "__main__":
    if not TEMPLATE.exists():
        raise SystemExit(f"No se encuentra la plantilla: {TEMPLATE}")
    for lang in ("ru", "es", "en"):
        print(f"Generando versión {lang}...")
        build(lang, CONTENT[lang])